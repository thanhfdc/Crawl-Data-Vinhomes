from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

DOCUMENT_HEADERS = ["No.", "Document ID", "Company ID", "Document Type", "Title", "Published Date", "Publisher", "Source Tier", "Source ID", "Source URL", "Canonical URL", "Parent Page URL", "Retrieved At", "Content Type", "File Size", "SHA256", "Storage Path", "Status"]
SOURCE_HEADERS = ["Source ID", "Company ID", "Source Name", "Source Type", "Source Tier", "Base URL", "Domain", "Enabled", "Last Crawl Status", "Last Checked At"]
JOB_HEADERS = ["Job ID", "Company ID", "Started At", "Finished At", "Status", "Sources Processed", "URLs Discovered", "Documents Downloaded", "Documents Skipped", "Documents Failed"]
ERROR_HEADERS = ["Timestamp", "Job ID", "Company ID", "Source ID", "URL", "Error Type", "HTTP Status", "Error Message", "Retry Attempt", "Duration Ms"]

def _sheet(wb, name, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    return ws

def _finish(ws, widths, wrap_columns=()):
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for column in wrap_columns:
        for cell in ws[column]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

def _link(cell, value):
    if value:
        cell.value = value
        cell.hyperlink = value
        cell.style = "Hyperlink"

def export_excel(store, config, output):
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    documents = store.export(config.company_id)
    ws = _sheet(wb, "Documents", DOCUMENT_HEADERS)
    for no, row in enumerate(documents, 1):
        values = [no, row.get("document_id"), row.get("company_id"), row.get("document_type"), row.get("title"), row.get("published_at"), row.get("publisher"), row.get("source_tier"), row.get("source_id"), row.get("source_url"), row.get("canonical_url"), row.get("parent_page_url"), row.get("retrieved_at"), row.get("content_type"), row.get("file_size"), row.get("sha256"), row.get("storage_path"), "stored"]
        ws.append(values)
        _link(ws.cell(ws.max_row, 10), row.get("source_url"))
        _link(ws.cell(ws.max_row, 11), row.get("canonical_url"))
        _link(ws.cell(ws.max_row, 12), row.get("parent_page_url"))
    _finish(ws, [7, 34, 12, 28, 55, 20, 18, 12, 28, 55, 55, 55, 25, 22, 14, 68, 55, 14], ("E", "J", "K", "L", "P", "Q"))

    source_ws = _sheet(wb, "Sources", SOURCE_HEADERS)
    for source in config.sources:
        source_ws.append([source.source_id, source.company_id, source.name, source.source_type, source.source_tier, source.base_url, source.domain, source.enabled, None, None])
        _link(source_ws.cell(source_ws.max_row, 6), source.base_url)
    _finish(source_ws, [28, 12, 34, 28, 12, 60, 28, 12, 20, 25], ("F",))
    _sheet(wb, "Crawl_Jobs", JOB_HEADERS)
    _sheet(wb, "Crawl_Errors", ERROR_HEADERS)
    wb.save(output)
    return {"path": str(output), "documents": len(documents), "sources": len(config.sources), "crawl_jobs": 0, "errors": 0}
