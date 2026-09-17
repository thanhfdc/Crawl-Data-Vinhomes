import shutil
from pathlib import Path
from openpyxl import load_workbook
from crawler.config import load_company
from crawler.excel_export import export_excel
from crawler.storage import Store

def test_excel_export():
    config = load_company("config/companies/vhm.yaml")
    tmp_path = Path(".test-artifacts/excel-export")
    shutil.rmtree(tmp_path, ignore_errors=True)
    tmp_path.mkdir(parents=True)
    store = Store(tmp_path / "data")
    try:
        result = export_excel(store, config, tmp_path / "out" / "vhm.xlsx")
        workbook = load_workbook(result["path"])
        assert workbook.sheetnames == ["Documents", "Sources", "Crawl_Jobs", "Crawl_Errors"]
        assert [cell.value for cell in workbook["Documents"][1]] == ["No.", "Document ID", "Company ID", "Document Type", "Title", "Published Date", "Publisher", "Source Tier", "Source ID", "Source URL", "Canonical URL", "Parent Page URL", "Retrieved At", "Content Type", "File Size", "SHA256", "Storage Path", "Status"]
        assert workbook["Documents"].max_row == 1
        assert workbook["Crawl_Jobs"].max_row == 1
        assert workbook["Crawl_Errors"].max_row == 1
    finally:
        store.close()
        shutil.rmtree(tmp_path, ignore_errors=True)
